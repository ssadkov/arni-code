# Copyright (c) Microsoft Corporation. All rights reserved.
# Licensed under the MIT License. See License.txt in the project root.
"""Behavioral checks for the bounded first-version profession helpers."""
import copy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]


def load(skill, filename):
    spec = importlib.util.spec_from_file_location(filename, ROOT / 'src/vs/sessions/skills' / skill / 'scripts' / (filename + '.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


tables = load('spreadsheet-analysis', 'tables')
slides = load('presentation-design', 'slides')
lesson = load('lesson-design', 'lesson')
organize = load('file-organizing', 'organize')


class SpecialistTest(unittest.TestCase):
    def setUp(self):
        base = ROOT / '.build/specialist-test-tmp'
        base.mkdir(parents=True, exist_ok=True)
        assert base.resolve().is_relative_to(ROOT)
        self.temp = tempfile.TemporaryDirectory(dir=base)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def file(self, name, text):
        path = self.root / name
        path.write_text(text, encoding='utf-8')
        return path

    def test_reconciliation_preserves_ambiguity_and_decimal_accuracy(self):
        left = self.file('left.csv', 'id,amount\n001,0.10\n002,200\n003,50\n005,1\n005,2\n006,unknown\n,2\n')
        right = self.file('right.csv', 'id,amount\n001,0.1\n002,210\n004,80\n005,3\n006,4\n')
        before = left.read_bytes(), right.read_bytes()
        result = tables.operate('reconcile', [left, right], self.root / 'result.xlsx', 'id', 'amount')
        self.assertEqual(result['statuses'], {'invalid_key': 1, 'matched': 1, 'amount_difference': 1, 'left_only': 1, 'right_only': 1, 'ambiguous': 1, 'invalid_amount': 1})
        rows = tables.read(self.root / 'result.xlsx')[1]
        self.assertEqual(rows[1]['id'], '001')
        self.assertEqual(rows[2]['difference'], '-10')
        self.assertEqual(before, (left.read_bytes(), right.read_bytes()))

    def test_merge_never_discards_duplicates_or_executes_formula_strings(self):
        a = self.file('a.csv', 'id,name\n01,A\n01,B\n,@note\n')
        b = self.file('b.csv', 'id,name\n02,C\n')
        out = self.root / 'merge.xlsx'
        self.assertEqual(tables.operate('merge', [a, b], out, 'id')['output_rows'], 4)
        rows = tables.read(out)[1]
        self.assertEqual([r['_key_status'] for r in rows], ['duplicate_key', 'duplicate_key', 'missing_key', 'unique'])
        with self.assertRaises(FileExistsError):
            tables.operate('merge', [a, b], out, 'id')
        tables.write(self.root / 'literal.xlsx', ['value'], [['=1+1']])
        from openpyxl import load_workbook
        book = load_workbook(self.root / 'literal.xlsx')
        self.assertEqual(book.active['A2'].data_type, 's')
        book.close()
        formula = self.file('formula.csv', 'id,amount\n01,=1+1\n')
        with self.assertRaisesRegex(ValueError, 'Formula'):
            tables.operate('merge', [formula], self.root / 'bad.csv', 'id')
        self.assertFalse((self.root / 'bad.csv').exists())

    def test_csv_output_neutralizes_active_cells(self):
        out = self.root / 'safe.csv'
        tables.write(out, ['name'], [['=HYPERLINK("x")'], ['+cmd'], ['-text'], ['@SUM(1)']])
        self.assertTrue(all(row['name'].startswith("'") for row in tables.read(out)[1]))

    def test_presentation_creation_and_narrow_edit_preserve_other_slides(self):
        from pptx import Presentation
        plan = {'slides': [{'title': 'First', 'bullets': ['Evidence'], 'source': 'source.md'}, {'title': 'Second', 'bullets': ['Unchanged']}]}
        path, changed = self.root / 'deck.pptx', self.root / 'changed.pptx'
        self.assertEqual(len(slides.create(plan, path)['slides']), 2)
        slides.replace(path, 'First', 'Updated', changed)
        old, new = Presentation(path), Presentation(changed)
        self.assertEqual(new.slides[0].shapes[0].text, 'Updated')
        self.assertEqual(old.slides[0].shapes[0].text_frame.paragraphs[0].font.size, new.slides[0].shapes[0].text_frame.paragraphs[0].font.size)
        self.assertEqual(old.slides[1]._element.xml, new.slides[1]._element.xml)
        with self.assertRaises(ValueError):
            slides.create({'slides': [{'title': 'x', 'bullets': ['x' * 151]}]}, self.root / 'overflow.pptx')
        with self.assertRaises(FileExistsError):
            slides.create(plan, path)

    def test_lesson_is_source_anchored_and_escapes_payload(self):
        passage = 'The source contains </script> and __TITLE__. Water freezes at zero.'
        plan = {'title': '__PASSAGE__', 'passage': passage, 'source': 'source.md', 'questions': [{'question': '__TITLE__?', 'choices': ['Yes', 'No'], 'answer': 0, 'explanation': 'Source says so.', 'evidence': passage}]}
        path = self.root / 'lesson.html'
        lesson.create(plan, path)
        page = path.read_text(encoding='utf-8')
        self.assertEqual(page.count('</script>'), 1)
        self.assertIn('<title>__PASSAGE__</title>', page)
        payload = page.split('const questions=', 1)[1].split(';const results=', 1)[0]
        self.assertEqual(json.loads(payload), plan['questions'])
        plan['questions'][0]['evidence'] = 'Not in source'
        with self.assertRaises(ValueError):
            lesson.create(plan, self.root / 'invalid.html')

    def test_file_plan_is_read_only_and_moves_can_be_undone(self):
        folder = self.root / 'files'
        folder.mkdir()
        (folder / 'a.txt').write_bytes(b'original')
        (folder / 'b.csv').write_bytes(b'id\n001')
        plan = organize.plan(folder, self.root / 'plan.json')
        self.assertEqual(sorted(p.name for p in folder.iterdir()), ['a.txt', 'b.csv'])
        journal = self.root / 'journal.json'
        self.assertEqual(organize.apply(plan, journal)['moved'], 2)
        self.assertEqual((folder / 'Documents/a.txt').read_bytes(), b'original')
        self.assertEqual(organize.undo(journal)['restored'], 2)
        self.assertEqual((folder / 'a.txt').read_bytes(), b'original')
        self.assertFalse((folder / 'Documents/a.txt').exists())
        self.assertEqual(organize.undo(journal)['restored'], 0)

    def test_stale_plan_and_collisions_never_overwrite(self):
        folder = self.root / 'files'
        folder.mkdir()
        original = folder / 'a.txt'
        original.write_text('old')
        plan = organize.plan(folder, self.root / 'plan.json')
        original.write_text('new')
        with self.assertRaisesRegex(ValueError, 'Stale'):
            organize.apply(plan, self.root / 'journal.json')
        self.assertFalse((self.root / 'journal.json').exists())
        (folder / 'Documents').mkdir()
        (folder / 'Documents/a.txt').write_text('keep')
        plan = organize.plan(folder, self.root / 'plan2.json')
        self.assertEqual(plan['moves'], [])
        self.assertEqual(original.read_text(), 'new')

    def test_untrusted_plan_and_interrupted_journal_refused(self):
        folder = self.root / 'files'
        folder.mkdir()
        (folder / 'a.txt').write_text('content')
        plan = organize.plan(folder, self.root / 'plan.json')
        for bad in ('../escape', str(self.root / 'escape'), 'Documents/../escape', 'C:escape'):
            data = copy.deepcopy(plan)
            data['moves'][0]['to'] = bad
            with self.assertRaises(ValueError):
                organize.apply(data, self.root / 'journal.json')
        (folder / 'Documents').mkdir()
        (folder / 'Documents/a.txt').write_text('partial copy')
        journal = self.root / 'interrupted.json'
        organize.dump_new(journal, dict(plan, moves=[dict(plan['moves'][0], status='pending')]))
        with self.assertRaisesRegex(ValueError, 'Interrupted'):
            organize.undo(journal)
        self.assertEqual((folder / 'a.txt').read_text(), 'content')


if __name__ == '__main__':
    unittest.main(verbosity=2)
