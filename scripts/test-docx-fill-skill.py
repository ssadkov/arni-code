# Copyright (c) Microsoft Corporation. All rights reserved.
# Licensed under the MIT License. See License.txt in the project root.
"""docx-fill helper: blank indexes, formatting, refusals, and PDF/PNG when a converter exists."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
HELPER = ROOT / 'src/vs/sessions/skills/docx-fill/scripts/docx_fill.py'


def run(*args, ok=True):
	result = subprocess.run([sys.executable, str(HELPER), *map(str, args)], capture_output=True, text=True, encoding='utf-8')
	data = json.loads(result.stdout)
	assert (result.returncode == 0) == ok, (args, result.returncode, data)
	return data


def main():
	with tempfile.TemporaryDirectory() as directory:
		work = Path(directory)
		template = work / 'template.docx'
		doc = Document()
		paragraph = doc.add_paragraph()
		paragraph.add_run('Стороны: ').bold = True
		paragraph.add_run('______, БИН ______ и ______, БИН ______.')
		doc.add_paragraph('Дата: «___» ________ 20__ г.')
		table = doc.add_table(rows=2, cols=2)
		table.cell(0, 0).text, table.cell(0, 1).text = 'ИИК', '______'
		table.cell(1, 0).text, table.cell(1, 1).text = 'Телефон', '______'
		for run_ in doc.paragraphs[0].runs:
			run_.font.size = Pt(13)
		doc.save(template)
		before = template.read_bytes()

		inspected = run('inspect', template)
		assert [item['at'] for item in inspected['items']] == ['p0', 'p1', 't0.r0.c1', 't0.r1.c1'], inspected['items']
		assert inspected['blank_count'] == 9 and inspected['tables'][0]['first_column'] == ['ИИК', 'Телефон']

		spec = work / 'spec.json'
		spec.write_text(json.dumps({'operations': [
			{'at': 'p0', 'nth': 1, 'replace': 'ТОО «Эдбиз»'},
			{'at': 'p0', 'nth': 4, 'replace': '190545260187'},
			{'at': 'p0', 'nth': 2, 'replace': '221040013762'},
			{'at': 'p1', 'set': 'Дата: «02» ноября 2026 г.'},
			{'at': 't0.r0.c1', 'set': 'KZ388562203125947874'},
		]}, ensure_ascii=False), encoding='utf-8')
		output = work / 'out' / 'filled.docx'
		filled = run('fill', '--template', template, '--spec', spec, '--output', output)
		result = Document(output)
		# nth counts blanks in the template, independent of the order of operations.
		assert result.paragraphs[0].text == 'Стороны: ТОО «Эдбиз», БИН 221040013762 и ______, БИН 190545260187.', result.paragraphs[0].text
		assert result.paragraphs[0].runs[0].bold and all(r.font.size == Pt(13) for r in result.paragraphs[0].runs if r.text)
		assert [b['at'] for b in filled['remaining_blanks']] == ['p0', 't0.r1.c1'] and filled['applied'] == 5
		assert template.read_bytes() == before, 'template must not change'

		run('fill', '--template', template, '--spec', spec, '--output', output, ok=False)
		run('fill', '--template', template, '--spec', spec, '--output', template, ok=False)
		mixed = work / 'mixed.json'
		mixed.write_text(json.dumps([{'at': 'p1', 'set': 'x'}, {'at': 'p1', 'replace': 'y'}]), encoding='utf-8')
		run('fill', '--template', template, '--spec', mixed, '--output', work / 'mixed.docx', ok=False)
		missing = work / 'missing.json'
		missing.write_text(json.dumps([{'at': 'p0', 'nth': 9, 'replace': 'z'}]), encoding='utf-8')
		assert run('fill', '--template', template, '--spec', missing, '--output', work / 'missing.docx', ok=False)['errors'][0]['error'] == 'match not found'

		doctor = subprocess.run([sys.executable, str(HELPER), 'doctor'], capture_output=True, text=True, encoding='utf-8')
		if json.loads(doctor.stdout)['pdf_converter']:
			pdf = run('pdf', output)
			assert pdf['pages'] == 1
			pages = run('render', pdf['pdf'], '--output-dir', work / 'pages')['pages']
			assert len(pages) == 1 and Path(pages[0]['png']).stat().st_size > 1000 and '221040013762' in pages[0]['text_preview']
			assert '221040013762' in run('text', pdf['pdf'])['text']
			print('PASS docx-fill with PDF and page images')
		else:
			print('PASS docx-fill; PDF skipped: no Word or LibreOffice')


if __name__ == '__main__':
	main()
