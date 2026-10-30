# Copyright (c) Microsoft Corporation. All rights reserved.
# Licensed under the MIT License. See License.txt in the project root.
"""Create the pptx-update test project: an outdated course deck and the brief it must be brought in line with."""
import hashlib
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

ACCENT = RGBColor(0x24, 0x57, 0xFF)
TITLE_COLOR = RGBColor(0x1D, 0x1D, 0x1F)

# Deliberately outdated: wrong start date, price, duration, schedule and e-mail.
SLIDES = [
	('AI-агенты для бизнеса', ['Курс Арнион', 'Старт 15.09.2026']),
	('Для кого', ['Владельцы малого бизнеса', 'Руководители отделов', 'Бухгалтеры и офис-менеджеры']),
	('Что вы получите', ['Три работающих сценария в Arni Code на своих данных', 'Навык ставить агенту задачу и проверять результат']),
	('Стоимость и формат', ['390 000 тенге', '6 недель, онлайн', 'Занятия по вторникам в 18:00']),
	('Контакты', ['hello@arnion.kz', 'arnion.kz']),
]

BRIEF = """# Бриф: обновление презентации курса

Актуальные факты. Всё, что в презентации им противоречит, устарело.

- Старт: 02.11.2026
- Длительность и формат: 4 недели, онлайн
- Расписание: понедельник и четверг в 19:00 (Астана)
- Стоимость: 450 000 тенге
- Контакты: e-mail и телефон организатора — со страницы https://arnion.kz/contacts/

## Программа (новый слайд после «Что вы получите»)

1. Что такое AI-агент и где он экономит время
2. Документы: договоры, счета и акты по реквизитам
3. Сайты и формы: лендинг за вечер
4. Свой агент для отдела: навыки, проверки, бюджет
"""


def build_deck(path: Path) -> None:
	deck = Presentation()
	deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
	layout = deck.slide_layouts[1]
	for title, bullets in SLIDES:
		slide = deck.slides.add_slide(layout)
		accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(228600), deck.slide_height)
		accent.name = 'Accent'
		accent.fill.solid()
		accent.fill.fore_color.rgb = ACCENT
		accent.line.fill.background()
		# Set the whole box: a partial xfrm leaves top/height at 0 and the title overlaps the body.
		slide.shapes.title.left, slide.shapes.title.top = Emu(609600), Emu(457200)
		slide.shapes.title.width, slide.shapes.title.height = Emu(10972800), Emu(1143000)
		slide.shapes.title.text = title
		for run in slide.shapes.title.text_frame.paragraphs[0].runs:
			run.font.size, run.font.bold, run.font.color.rgb = Pt(36), True, TITLE_COLOR
		body = slide.placeholders[1]
		body.left, body.top, body.width, body.height = Emu(609600), Emu(1828800), Emu(10972800), Emu(4572000)
		body.text_frame.text = bullets[0]
		for bullet in bullets[1:]:
			body.text_frame.add_paragraph().text = bullet
		for paragraph in body.text_frame.paragraphs:
			for run in paragraph.runs:
				run.font.size = Pt(24)
	deck.save(path)


def main() -> None:
	target = Path(sys.argv[1] if len(sys.argv) > 1 else 'C:/work/test-arni-pptx').resolve()
	if (target / 'source').exists():
		sys.exit(f'{target} already contains source/; remove it or choose another folder')
	(target / 'source').mkdir(parents=True)
	deck = target / 'source' / 'Курс_AI-агенты.pptx'
	build_deck(deck)
	(target / 'brief.md').write_text(BRIEF, encoding='utf-8')
	manifest = {'deck': deck.name, 'sha256': hashlib.sha256(deck.read_bytes()).hexdigest()}
	(target / '.case.json').write_text(json.dumps(manifest, ensure_ascii=False, indent='\t'), encoding='utf-8')
	print(f'Created {target}')


if __name__ == '__main__':
	main()
