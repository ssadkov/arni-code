# Copyright (c) Microsoft Corporation. All rights reserved.
# Licensed under the MIT License. See License.txt in the project root.
"""Verify the updated deck: current facts, new program slide in the deck's style, shapes on the slide, previews rendered."""
import hashlib
import json
from pathlib import Path
import re
import sys

from pptx import Presentation

MODULES = [
	'Что такое AI-агент и где он экономит время',
	'Документы: договоры, счета и акты по реквизитам',
	'Сайты и формы: лендинг за вечер',
	'Свой агент для отдела: навыки, проверки, бюджет',
]
OUTDATED = ['15.09.2026', '390 000', '6 недель', 'вторник', '18:00', 'hello@arnion.kz']
TITLES = ['AI-агенты для бизнеса', 'Для кого', 'Что вы получите', 'Программа', 'Стоимость и формат', 'Контакты']


def norm(text: str) -> str:
	return re.sub(r'\s+', ' ', text.replace('\u00a0', ' ').replace('\u202f', ' ')).strip()


def slide_text(slide) -> str:
	return norm(' '.join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame))


def accent(slide):
	return next((shape for shape in slide.shapes if shape.shape_type == 1 and shape.left <= 0 and shape.width < 400000
		and shape.fill.type == 1 and str(shape.fill.fore_color.rgb) == '2457FF'), None)


def main() -> int:
	project = Path(sys.argv[1] if len(sys.argv) > 1 else 'C:/work/test-arni-pptx').resolve()
	results = []

	def check(name: str, ok: bool, detail: str = '') -> None:
		results.append({'check': name, 'ok': bool(ok), **({'detail': detail[:200]} if detail and not ok else {})})

	case = json.loads((project / '.case.json').read_text(encoding='utf-8'))
	source = project / 'source' / case['deck']
	check('исходная презентация не изменена', hashlib.sha256(source.read_bytes()).hexdigest() == case['sha256'])
	decks = sorted((project / 'output').glob('*.pptx')) if (project / 'output').exists() else []
	check('ровно один PPTX в output/', len(decks) == 1, ', '.join(path.name for path in decks))
	if len(decks) == 1:
		deck = Presentation(decks[0])
		slides = list(deck.slides)
		titles = [norm(slide.shapes.title.text) if slide.shapes.title else '' for slide in slides]
		check('порядок слайдов', titles == TITLES, ' | '.join(titles))
		text = ' '.join(slide_text(slide) for slide in slides)
		found = [value for value in OUTDATED if value in text]
		check('не осталось устаревших фактов', not found, ', '.join(found))
		check('дата старта', '02.11.2026' in text)
		check('стоимость', '450 000' in text and 'тенге' in text)
		check('длительность', '4 недели' in text)
		check('расписание', '19:00' in text and re.search(r'понедельник|пн', text, re.I) and re.search(r'четверг|чт', text, re.I))
		check('e-mail с сайта', 'zakaz@arnion.kz' in text)
		check('телефон с сайта', '77057252942' in re.sub(r'\D', '', text))
		program = slides[3] if len(slides) > 3 else None
		program_text = slide_text(program) if program else ''
		positions = [program_text.find(title) for title in MODULES]
		check('программа: 4 модуля дословно и по порядку', all(position >= 0 for position in positions) and positions == sorted(positions),
			program_text)
		missing = [index + 1 for index, slide in enumerate(slides) if not accent(slide)]
		check('синяя полоса стиля на каждом слайде', not missing, f'нет на слайдах {missing}')
		if program and slides:
			original = slides[1].shapes.title
			title = program.shapes.title
			check('заголовок нового слайда в стиле колоды', title is not None and abs(title.left - original.left) < 120000
				and abs(title.top - original.top) < 120000, f'{title and (title.left, title.top)} vs {(original.left, original.top)}')
		outside = [f'{index + 1}:{shape.name}' for index, slide in enumerate(slides) for shape in slide.shapes
			if shape.left is not None and (shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width + 9525
				or shape.top + shape.height > deck.slide_height + 9525)]
		check('фигуры не выходят за слайд', not outside, ', '.join(outside))
		def overlaps(a, b) -> bool:
			return a.left < b.left + b.width and b.left < a.left + a.width and a.top < b.top + b.height and b.top < a.top + a.height
		crowded = []
		for index, slide in enumerate(slides):
			boxes = [shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]
			if any(overlaps(first, second) for position, first in enumerate(boxes) for second in boxes[position + 1:]):
				crowded.append(index + 1)
		check('текстовые блоки не наезжают друг на друга', not crowded, f'слайды {sorted(set(crowded))}')
		small = [f'{index + 1}' for index, slide in enumerate(slides) for shape in slide.shapes if shape.has_text_frame
			for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size and run.font.size.pt < 14]
		check('текст не мельче 14 pt', not small, f'слайды {sorted(set(small))}')
		previews = sorted((project / 'output' / 'preview').glob('*.png')) if (project / 'output' / 'preview').exists() else []
		check('превью каждого слайда', len(previews) == len(slides), f'{len(previews)} PNG на {len(slides)} слайдов')

	failed = [result for result in results if not result['ok']]
	sys.stdout.reconfigure(encoding='utf-8')
	print(json.dumps({'passed': len(results) - len(failed), 'total': len(results), 'failed': failed}, ensure_ascii=False, indent='\t'))
	return 1 if failed else 0


if __name__ == '__main__':
	sys.exit(main())
