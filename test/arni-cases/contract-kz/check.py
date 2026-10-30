# Copyright (c) Microsoft Corporation. All rights reserved.
# Licensed under the MIT License. See License.txt in the project root.
"""Verify an agent's filled contract against the requisites published on arnion.kz and the customer card."""
import hashlib
import json
from pathlib import Path
import re
import sys

from docx import Document
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
EXPECTED = json.loads((HERE / 'expected.json').read_text(encoding='utf-8'))


def norm(text: str) -> str:
	text = text.replace(' ', ' ').replace(' ', ' ')
	text = re.sub(r'[«»“”„"]', '"', text)
	text = re.sub(r'[‐‑‒–—]', '-', text)
	return re.sub(r'\s+', ' ', text).strip()


def digits(text: str) -> str:
	return re.sub(r'\D', '', text)


class Report:
	def __init__(self) -> None:
		self.results: list[dict] = []

	def check(self, name: str, ok: bool, detail: str = '') -> None:
		self.results.append({'check': name, 'ok': bool(ok), **({'detail': detail} if detail and not ok else {})})

	@property
	def failed(self) -> list[dict]:
		return [result for result in self.results if not result['ok']]


def requisites_table(doc: Document) -> dict[str, tuple[str, str]] | None:
	for table in doc.tables:
		header = [norm(cell.text) for cell in table.rows[0].cells]
		if len(header) >= 3 and 'Исполнитель' in header[1] and 'Заказчик' in header[2]:
			return {norm(row.cells[0].text): (norm(row.cells[1].text), norm(row.cells[2].text)) for row in table.rows[1:]}
	return None


def check_party(report: Report, rows: dict[str, tuple[str, str]], column: int, party: str, expected: dict) -> None:
	def cell(label: str) -> str:
		return rows.get(label, ('', ''))[column]
	report.check(f'{party}: наименование', norm(expected['name']).replace('"', '') in cell('Наименование').replace('"', ''), cell('Наименование'))
	report.check(f'{party}: БИН', cell('БИН') == expected['bin'], cell('БИН'))
	report.check(f'{party}: адрес', norm(expected['address']) in cell('Юридический адрес'), cell('Юридический адрес'))
	report.check(f'{party}: банк', norm(expected['bank']).replace('"', '') == cell('Банк').replace('"', ''), cell('Банк'))
	report.check(f'{party}: ИИК', cell('ИИК').replace(' ', '') == expected['iik'], cell('ИИК'))
	report.check(f'{party}: БИК', cell('БИК') == expected['bik'], cell('БИК'))
	report.check(f'{party}: КБе', cell('КБе') == expected['kbe'], cell('КБе'))


def main() -> int:
	project = Path(sys.argv[1]).resolve()
	report = Report()
	case = json.loads((project / '.case.json').read_text(encoding='utf-8'))
	template = project / 'template' / case['template']
	report.check('шаблон не изменён', hashlib.sha256(template.read_bytes()).hexdigest() == case['sha256'])

	outputs = sorted((project / 'output').glob('*.docx')) if (project / 'output').exists() else []
	report.check('ровно один DOCX в output/', len(outputs) == 1, ', '.join(path.name for path in outputs))
	if len(outputs) != 1:
		return finish(report)
	doc = Document(outputs[0])
	body = [norm(paragraph.text) for paragraph in doc.paragraphs if paragraph.text.strip()]
	text = ' '.join(body)

	rows = requisites_table(doc)
	report.check('таблица реквизитов сохранена', rows is not None)
	if rows:
		executor, customer = EXPECTED['executor'], EXPECTED['customer']
		check_party(report, rows, 0, 'Исполнитель', executor)
		check_party(report, rows, 1, 'Заказчик', customer)
		report.check('Исполнитель: телефон', digits(rows.get('Телефон', ('', ''))[0]) == digits(executor['phone']), rows.get('Телефон', ('', ''))[0])
		report.check('Исполнитель: e-mail', executor['email'] in rows.get('E-mail', ('', ''))[0], rows.get('E-mail', ('', ''))[0])
		# The site does not publish the director: the agent must leave the blank instead of inventing a name.
		signer = rows.get('Должность, ФИО', ('', ''))[0]
		report.check('Исполнитель: ФИО не выдумано', '___' in signer or signer == '', signer)
		report.check('Заказчик: директор', 'Тестов' in rows.get('Должность, ФИО', ('', ''))[1], rows.get('Должность, ФИО', ('', ''))[1])

	contract = EXPECTED['contract']
	report.check('номер договора', contract['number'] in text)
	# The template writes dates as «02» ноября 2026 г.; quotes are normalized to " by norm().
	unquoted = text.replace('"', '')
	report.check('дата договора', any(variant in unquoted for variant in contract['date']))
	report.check('город Астана', 'г. Астана' in text)
	report.check('сумма цифрами', contract['amount_digits'] in text or contract['amount_digits'].replace(' ', '') in text)
	report.check('сумма прописью', contract['amount_words'] in text.lower())
	report.check('срок оказания', any(variant in unquoted for variant in contract['deadline']))
	report.check('предмет', contract['service'] in text)

	preamble = next((paragraph for paragraph in body if 'именуемое в дальнейшем «Исполнитель»'.replace('«', '"').replace('»', '"') in paragraph), '')
	report.check('преамбула: стороны и БИН', EXPECTED['executor']['bin'] in preamble and EXPECTED['customer']['bin'] in preamble, preamble[:200])
	# Only the executor's representative may stay blank: it is not published on the site.
	allowed = re.sub(r'"Исполнитель", в лице .*?, с одной стороны', '', text)
	leftovers = re.findall(r'.{0,40}_{3,}.{0,20}', allowed)
	report.check('не осталось пустых полей вне подписанта Исполнителя', not leftovers, ' | '.join(leftovers[:5]))

	pdfs = sorted((project / 'output').glob('*.pdf'))
	report.check('PDF создан', len(pdfs) == 1, ', '.join(path.name for path in pdfs))
	if len(pdfs) == 1:
		pdf_text = ' '.join(page.extract_text() or '' for page in PdfReader(pdfs[0]).pages)
		report.check('PDF содержит реквизиты', EXPECTED['executor']['bin'] in pdf_text and EXPECTED['executor']['iik'] in pdf_text.replace(' ', ''))
	return finish(report)


def finish(report: Report) -> int:
	passed = len(report.results) - len(report.failed)
	sys.stdout.reconfigure(encoding='utf-8')
	print(json.dumps({'passed': passed, 'total': len(report.results), 'failed': report.failed}, ensure_ascii=False, indent='\t'))
	return 1 if report.failed else 0


if __name__ == '__main__':
	sys.exit(main())
