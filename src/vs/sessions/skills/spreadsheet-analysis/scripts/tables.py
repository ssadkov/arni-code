#!/usr/bin/env python3
# Copyright (c) Microsoft Corporation. Licensed under the MIT License.
"""Bounded CSV/XLSX inspection, merge and reconciliation. Never overwrites input."""
import argparse
from collections import Counter, defaultdict
import csv
import importlib.util
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import sys
import subprocess

sys.dont_write_bytecode = True

LIMIT = 50000


def read(path, sheet=None, delimiter=','):
    path = Path(path)
    if path.stat().st_size > 32 * 1024**2:
        raise ValueError('Input exceeds 32 MiB')
    if path.suffix.lower() in ('.csv', '.tsv'):
        with path.open(encoding='utf-8-sig', newline='') as handle:
            rows = []
            for row in csv.reader(handle, delimiter='\t' if path.suffix.lower() == '.tsv' else delimiter):
                rows.append(row)
                if len(rows) > LIMIT + 1:
                    raise ValueError('Row limit exceeded')
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            page = book[sheet] if sheet else book.worksheets[0]
            if (page.max_row or 0) > LIMIT + 1 or (page.max_column or 0) > 200:
                raise ValueError('Worksheet dimensions exceed row/column limits')
            rows = []
            for row in page.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > LIMIT + 1:
                    raise ValueError('Row limit exceeded')
        finally:
            book.close()
    else:
        raise ValueError('Use CSV, TSV or XLSX; legacy XLS/macros are unsupported')
    if not rows or len(rows[0]) > 200:
        raise ValueError('Empty table or too many columns')
    headers = [str(value).strip() if value is not None else '' for value in rows[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError('Headers must be unique and nonempty')
    result = []
    for row in rows[1:]:
        if not any(value is not None and value != '' for value in row):
            continue
        if len(row) > len(headers):
            raise ValueError('Row has more cells than headers')
        result.append(dict(zip(headers, row + [None] * (len(headers) - len(row)))))
    return headers, result


def formula(value):
    return isinstance(value, str) and value.startswith('=')


def key(value):
    return '' if value is None else str(value).strip()


def amount(value):
    if isinstance(value, bool) or value is None or ',' in str(value):
        raise ValueError('Amount must use an explicit decimal point')
    try:
        result = Decimal(str(value).strip())
    except InvalidOperation as error:
        raise ValueError('Invalid amount') from error
    if not result.is_finite():
        raise ValueError('Non-finite amount')
    return result


def write(path, headers, records):
    path = Path(path)
    if path.exists():
        raise FileExistsError('Choose a new output file')
    if path.suffix.lower() == '.xlsx':
        from openpyxl import Workbook
        book = Workbook()
        page = book.active
        page.title = 'Results'
        for index, values in enumerate([headers, *records], 1):
            page.append(values)
            for cell in page[index]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        page.freeze_panes = 'A2'
        page.auto_filter.ref = page.dimensions
        with path.open('xb') as handle:
            book.save(handle)
        book.close()
    elif path.suffix.lower() == '.csv':
        with path.open('x', encoding='utf-8-sig', newline='') as handle:
            # Neutralize active cells on opening CSV in a spreadsheet application.
            safe = lambda value: "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r')) else value
            writer = csv.writer(handle)
            writer.writerow([safe(value) for value in headers])
            writer.writerows([[safe(value) for value in row] for row in records])
    else:
        raise ValueError('Output must be CSV or XLSX')


def operate(command, paths, output=None, key_column=None, amount_column=None, sheet=None, delimiter=','):
    tables = [read(path, sheet, delimiter) for path in paths]
    if command == 'inspect':
        headers, rows = tables[0]
        return {'columns': headers, 'rows': len(rows), 'formula_cells': sum(formula(value) for row in rows for value in row.values()), 'sample': rows[:3], 'formula_recalculation': False}
    if any(formula(value) for _, rows in tables for row in rows for value in row.values()):
        raise ValueError('Formula input requires a recalculated values export')
    if any(key_column not in headers for headers, _ in tables):
        raise ValueError('Key column is absent')
    records = []
    if command == 'merge':
        headers = tables[0][0]
        if {'_source', '_key_status'} & set(headers):
            raise ValueError('Input uses reserved merge metadata columns')
        if any(columns != headers for columns, _ in tables):
            raise ValueError('Merge requires identical ordered headers')
        counts = Counter(key(row[key_column]) for _, rows in tables for row in rows)
        for path, (_, rows) in zip(paths, tables):
            for row in rows:
                value = key(row[key_column])
                records.append([Path(path).name, 'missing_key' if not value else 'duplicate_key' if counts[value] > 1 else 'unique', *row.values()])
        write(output, ['_source', '_key_status', *headers], records)
        return {'input_rows': sum(len(rows) for _, rows in tables), 'output_rows': len(records), 'duplicates_preserved': True}
    if len(tables) != 2 or any(amount_column not in headers for headers, _ in tables):
        raise ValueError('Reconciliation requires two tables and an amount column')
    groups = []
    for _, rows in tables:
        grouped = defaultdict(list)
        for row in rows:
            grouped[key(row[key_column])].append(row)
        groups.append(grouped)
    for identity in sorted(groups[0].keys() | groups[1].keys()):
        left, right = groups[0].get(identity, []), groups[1].get(identity, [])
        status = 'invalid_key' if not identity else 'ambiguous' if len(left) > 1 or len(right) > 1 else 'left_only' if not right else 'right_only' if not left else 'matched'
        a = b = delta = ''
        try:
            a = str(amount(left[0][amount_column])) if len(left) == 1 else ''
            b = str(amount(right[0][amount_column])) if len(right) == 1 else ''
            if status == 'matched':
                delta = str(Decimal(a) - Decimal(b))
                status = 'matched' if Decimal(delta) == 0 else 'amount_difference'
        except ValueError:
            status = status if status in ('ambiguous', 'invalid_key') else 'invalid_amount'
        records.append([identity, status, len(left), len(right), a, b, delta])
    write(output, [key_column, 'status', 'left_rows', 'right_rows', 'left_amount', 'right_amount', 'difference'], records)
    return {'input_rows': [len(rows) for _, rows in tables], 'keys': len(records), 'statuses': dict(Counter(row[1] for row in records)), 'output': str(output)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['prepare', 'inspect', 'merge', 'reconcile'])
    parser.add_argument('paths', nargs='*')
    parser.add_argument('--key')
    parser.add_argument('--amount')
    parser.add_argument('--sheet')
    parser.add_argument('--delimiter', default=',')
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    if args.command == 'prepare':
        if args.paths or args.output or args.key or args.amount or args.sheet:
            parser.error('Prepare takes no input files or output options')
    else:
        if not args.paths:
            parser.error('At least one input file is required')
        if args.command != 'inspect' and (not args.key or not args.output):
            parser.error('--key and --output required')
        if args.command == 'reconcile' and not args.amount:
            parser.error('--amount required for reconciliation')
        if args.command == 'inspect' and len(args.paths) != 1:
            parser.error('Inspect one file at a time')
    if args.command == 'prepare' or any(Path(path).suffix.lower() == '.xlsx' for path in [*args.paths, str(args.output or '')]):
        # Load by absolute path so this also works with the managed interpreter's -I mode.
        spec = importlib.util.spec_from_file_location('arni_xlsx_runtime', Path(__file__).with_name('xlsx_runtime.py'))
        runtime = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(runtime)
        executable = runtime.prepare()
        if args.command == 'prepare':
            print(json.dumps({'status': 'ready', 'python_executable': str(executable)}, ensure_ascii=True))
            return
        if Path(sys.prefix).resolve() != executable.parent.parent.resolve():
            sys.exit(subprocess.run([str(executable), '-I', str(Path(__file__).resolve()), *sys.argv[1:]]).returncode)
    print(json.dumps(operate(args.command, args.paths, args.output, args.key, args.amount, args.sheet, args.delimiter), ensure_ascii=True, default=str))


if __name__ == '__main__':
    try:
        main()
    except (OSError, ValueError, ImportError, KeyError, RuntimeError) as error:
        print(json.dumps({'status': 'failed', 'reason': str(error)}, ensure_ascii=True), file=sys.stderr)
        sys.exit(2)
