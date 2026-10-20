#!/usr/bin/env python3
# Copyright (c) Microsoft Corporation. Licensed under the MIT License.
"""Simple editable PPTX and narrow format-preserving text replacement."""
import argparse
import json
from pathlib import Path
import sys


def inspect(path):
    from pptx import Presentation
    book = Presentation(path)
    return {'slides': [{'number': i + 1, 'text': [shape.text for shape in slide.shapes if shape.has_text_frame]} for i, slide in enumerate(book.slides)], 'visual_check': 'not_performed'}


def save(book, output):
    with Path(output).open('xb') as handle:
        book.save(handle)
    return inspect(output)


def create(plan, output):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    slides = plan.get('slides', [])
    if not 1 <= len(slides) <= 30:
        raise ValueError('Use 1 to 30 slides')
    # Validate all content before creating the output.
    for data in slides:
        if not isinstance(data.get('title'), str) or not 1 <= len(data['title']) <= 90:
            raise ValueError('A short title is required')
        bullets = data.get('bullets', [])
        if not isinstance(bullets, list) or len(bullets) > 6 or any(not isinstance(s, str) or len(s) > 150 for s in bullets):
            raise ValueError('Use up to six bullets of at most 150 characters')
        if not isinstance(data.get('source', ''), str) or len(data.get('source', '')) > 180:
            raise ValueError('Source footer is too long')
    book = Presentation()
    book.slide_width, book.slide_height = Inches(13.333), Inches(7.5)
    for i, data in enumerate(slides):
        slide = book.slides.add_slide(book.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(.7), Inches(.55), Inches(11.9), Inches(1.25)).text_frame
        title.word_wrap = True
        title.paragraphs[0].text = data['title']
        title.paragraphs[0].font.size = Pt(30)
        title.paragraphs[0].font.bold = True
        body = slide.shapes.add_textbox(Inches(.8), Inches(2), Inches(11.6), Inches(4.3)).text_frame
        body.word_wrap = True
        for j, text in enumerate(data.get('bullets', [])):
            paragraph = body.paragraphs[0] if j == 0 else body.add_paragraph()
            paragraph.text = text
            paragraph.font.size = Pt(21)
            paragraph.space_after = Pt(15)
        footer = slide.shapes.add_textbox(Inches(.8), Inches(6.7), Inches(11.6), Inches(.55)).text_frame
        footer.word_wrap = True
        footer.paragraphs[0].text = str(i + 1) + ' | ' + data.get('source', '')
        footer.paragraphs[0].font.size = Pt(10)
        footer.paragraphs[0].font.color.rgb = RGBColor(90, 90, 90)
    return save(book, output)


def replace(source, old, new, output):
    from pptx import Presentation
    if not old or not new or len(new) > 150 or '\n' in new:
        raise ValueError('Use a short nonempty single-line replacement')
    book = Presentation(source)
    matches = [shape for slide in book.slides for shape in slide.shapes if shape.has_text_frame and shape.text == old]
    if len(matches) != 1:
        raise ValueError('Text must match exactly one shape')
    paragraphs = matches[0].text_frame.paragraphs
    if len(paragraphs) != 1 or len(paragraphs[0].runs) != 1:
        raise ValueError('Rich-text replacement requires a separate copy-based edit')
    paragraphs[0].runs[0].text = new
    return save(book, output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['create', 'inspect', 'replace-text'])
    parser.add_argument('source', nargs='?')
    parser.add_argument('--plan', type=Path)
    parser.add_argument('--old')
    parser.add_argument('--new')
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    if args.command == 'create':
        if not args.plan or not args.output:
            parser.error('--plan and --output required')
        result = create(json.loads(args.plan.read_text(encoding='utf-8-sig')), args.output)
    elif args.command == 'replace-text':
        if not args.source or not args.old or not args.new or not args.output:
            parser.error('source, --old, --new and --output required')
        result = replace(args.source, args.old, args.new, args.output)
    else:
        if not args.source:
            parser.error('source required')
        result = inspect(args.source)
    print(json.dumps(result, ensure_ascii=True))


if __name__ == '__main__':
    try:
        main()
    except (OSError, ValueError, ImportError, KeyError) as error:
        print(json.dumps({'status': 'failed', 'reason': str(error)}, ensure_ascii=True), file=sys.stderr)
        sys.exit(2)
